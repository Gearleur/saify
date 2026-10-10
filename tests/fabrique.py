"""Fabrique de jeux de test : fichiers du 25/09 altérés en mémoire, base copiée et modifiée."""

import io
import shutil
import sqlite3
from pathlib import Path

import openpyxl

from reco import rapprocher

DONNEES = Path(__file__).resolve().parents[1] / "data" / "donnees"
PMS, A, B, C = (
    "positions_internes.csv",
    "depositaire_A_20260925.xlsx",
    "depositaire_B_20260925.txt",
    "depositaire_C_20260925.xml",
)
CLES = {"pms": PMS, "a": A, "b": B, "c": C}


def fichiers(**remplacements: bytes) -> dict[str, bytes]:
    tous = {nom: (DONNEES / nom).read_bytes() for nom in (PMS, A, B, C)}
    for cle, contenu in remplacements.items():
        tous[CLES[cle]] = contenu
    return tous


def lancer(
    depot: dict[str, bytes],
    db: Path | None = DONNEES / "orsay.db",
    param: Path = DONNEES / "parametrage.xlsx",
):
    return rapprocher(depot, param, db)


def regles(resultat, niveau: str) -> set[str]:
    return {a.regle for a in resultat.journal if a.niveau == niveau}


def remplacer(nom: str, *paires: tuple[str, str]) -> bytes:
    texte = (DONNEES / nom).read_bytes().decode()
    for avant, apres in paires:
        assert avant in texte, avant
        texte = texte.replace(avant, apres)
    return texte.encode()


def xlsx_a(
    titre: str | None = None,
    quantite: str | None = None,
    dupliquer: int | None = None,
    ajouter: list | None = None,
) -> bytes:
    wb = openpyxl.load_workbook(DONNEES / A)
    ws = wb.worksheets[0]
    if titre is not None:
        ws["A1"] = titre
    if quantite is not None:
        ws["D4"] = quantite
    if dupliquer is not None:
        ws.append([cellule.value for cellule in ws[dupliquer]])
    if ajouter is not None:
        ws.append(ajouter)
    sortie = io.BytesIO()
    wb.save(sortie)
    return sortie.getvalue()


def db_modifiee(dossier: Path, *requetes: str) -> Path:
    """Copie de orsay.db modifiée : l'original n'est jamais touché."""
    copie = dossier / "orsay.db"
    shutil.copy(DONNEES / "orsay.db", copie)
    with sqlite3.connect(copie) as db:
        for requete in requetes:
            db.execute(requete)
    return copie


def parametrage_modifie(dossier: Path, feuille: str, modifier) -> Path:
    """Copie de parametrage.xlsx dont la feuille est passée à `modifier(ws)`."""
    wb = openpyxl.load_workbook(DONNEES / "parametrage.xlsx")
    modifier(wb[feuille])
    copie = dossier / "parametrage.xlsx"
    wb.save(copie)
    return copie
