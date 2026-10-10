"""Préparation exploratoire des données Orsay : lecture, nettoyage tracé, export CSV.

Usage : uv run python exploration/prepare_data.py
Sorties dans exploration/clean/ :
- positions_unifiees.csv : toutes les lignes PMS + conservateurs A/B/C au même format
- journal_nettoyage.csv   : chaque transformation (fichier, ligne, brut, retenu, règle, niveau)
- comparatif.csv          : une ligne par (fonds, ISIN), quantités PMS / A / B / C côte à côte
- db_<table>.csv          : tables de orsay.db (lecture seule)
"""

import csv
import re
import sqlite3
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "donnees"
OUT = Path(__file__).resolve().parent / "clean"
DATE_ARRETE = "2026-09-25"
SOUS_COMPTES_EXCLUS = {"P3G"}


@dataclass
class Position:
    source: str  # PMS, A, B, C
    fichier: str
    ligne: int
    compte_brut: str
    fonds: str  # PAT / OBL / ACT ou "" si inconnu
    isin: str
    libelle: str
    conservateur: str  # A / B / C
    quantite: Decimal
    statut: str = "OK"  # OK, EXCLU, SIGNALE


@dataclass
class Trace:
    fichier: str
    ligne: int
    champ: str
    brut: str
    retenu: str
    niveau: str  # NORMALISE, EXCLU, SIGNALE, BLOQUANT
    regle: str


journal: list[Trace] = []


def tracer(fichier, ligne, champ, brut, retenu, niveau, regle):
    journal.append(Trace(fichier, ligne, champ, str(brut), str(retenu), niveau, regle))


# --- référentiels -------------------------------------------------------------


def lire_parametrage() -> dict[str, dict[str, str]]:
    wb = openpyxl.load_workbook(DATA / "parametrage.xlsx", read_only=True)
    feuilles = {}
    for ws in wb:
        lignes = list(ws.iter_rows(values_only=True))
        feuilles[ws.title] = [dict(zip(lignes[0], r, strict=True)) for r in lignes[1:]]
    return feuilles


def lire_db() -> dict[str, list[dict]]:
    con = sqlite3.connect(f"file:{DATA / 'orsay.db'}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    tables = [r[0] for r in con.execute("select name from sqlite_master where type='table'")]
    out = {t: [dict(r) for r in con.execute(f"select * from {t}")] for t in tables}
    con.close()
    return out


def isin_valide(isin: str) -> bool:
    """Contrôle de forme + clé de Luhn (lettres converties en 10..35)."""
    if not re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", isin):
        return False
    chiffres = "".join(str(int(c, 36)) for c in isin)
    total = 0
    for i, c in enumerate(reversed(chiffres)):
        n = int(c) * (2 if i % 2 else 1)
        total += n - 9 if n > 9 else n
    return total % 10 == 0


def conservateur_depuis_libelle(libelle: str) -> str | None:
    lib = libelle.upper()
    if "CACEIS" in lib:
        return "A"
    if "EUROCLEAR" in lib:
        return "B"
    if "BNP" in lib or "BP2S" in lib:
        return "C"
    return None


def normaliser_isin(fichier, ligne, brut: str) -> str:
    isin = brut.strip().upper().replace(" ", "")
    if isin != brut:
        tracer(fichier, ligne, "isin", brut, isin, "NORMALISE", "casse et espaces")
    return isin


def parse_quantite(fichier, ligne, brut: str, regle: str, transfo) -> Decimal | None:
    try:
        q = transfo(brut)
    except (InvalidOperation, ValueError):
        tracer(fichier, ligne, "quantite", brut, "", "BLOQUANT", "quantité illisible")
        return None
    if str(q) != brut:
        tracer(fichier, ligne, "quantite", brut, q, "NORMALISE", regle)
    return q


# --- lecteurs par source -----------------------------------------------------


def lire_pms(sc_vers_fonds: dict[str, str]) -> list[Position]:
    f = "positions_internes.csv"
    out = []
    with open(DATA / f, newline="", encoding="utf-8") as fh:
        for n, row in enumerate(csv.DictReader(fh, delimiter=";"), start=2):
            if row["date_arrete"] != DATE_ARRETE:
                tracer(f, n, "date_arrete", row["date_arrete"], "", "BLOQUANT", "date ≠ arrêté")
            isin = normaliser_isin(f, n, row["isin"])
            cons = conservateur_depuis_libelle(row["depositaire"])
            if cons is None:
                tracer(f, n, "depositaire", row["depositaire"], "", "BLOQUANT", "libellé inconnu")
            elif row["depositaire"] not in ("A", "B", "C"):
                tracer(
                    f, n, "depositaire", row["depositaire"], cons, "NORMALISE", "libellé → A/B/C"
                )
            q = parse_quantite(f, n, row["quantite"], "", Decimal)
            sc = row["sous_compte"].strip()
            statut, fonds = "OK", sc_vers_fonds.get(sc, "")
            if not sc:
                statut = "EXCLU"
                tracer(
                    f, n, "sous_compte", "", "", "EXCLU", "ligne sans sous-compte (anomalie saisie)"
                )
            elif sc in SOUS_COMPTES_EXCLUS:
                statut = "EXCLU"
                tracer(f, n, "sous_compte", sc, "", "EXCLU", "sous-compte P3G exclu")
            elif not fonds:
                tracer(f, n, "sous_compte", sc, "", "BLOQUANT", "sous-compte absent de la transco")
            out.append(
                Position(
                    "PMS", f, n, sc, fonds, isin, row["libelle"].strip(), cons or "", q, statut
                )
            )
    return out


def lire_a(compte_vers_fonds: dict[str, str]) -> list[Position]:
    f = "depositaire_A_20260925.xlsx"
    ws = openpyxl.load_workbook(DATA / f, read_only=True).active
    lignes = list(ws.iter_rows(values_only=True))
    titre = lignes[0][0] or ""
    if "25/09/2026" not in titre:
        tracer(f, 1, "titre", titre, "", "BLOQUANT", "date ≠ arrêté")
    entete = next(i for i, r in enumerate(lignes) if r[0] == "Compte")
    tracer(
        f,
        1,
        "entete",
        f"{entete} lignes de titre",
        f"en-tête ligne {entete + 1}",
        "NORMALISE",
        "sauter le titre",
    )
    out = []
    for i, (compte, isin, lib, qte) in enumerate(lignes[entete + 1 :], start=entete + 2):
        if compte is None:
            continue
        compte = str(compte)
        isin = normaliser_isin(f, i, str(isin))
        q = parse_quantite(
            f,
            i,
            str(qte),
            "séparateur de milliers",
            lambda s: Decimal(s.replace(" ", "").replace(" ", "")),
        )
        fonds = compte_vers_fonds.get(compte, "")
        if not fonds:
            tracer(f, i, "compte", compte, "", "BLOQUANT", "compte absent de la transco")
        out.append(Position("A", f, i, compte, fonds, isin, str(lib).strip(), "A", q))
    return out


def lire_b(desk_vers_fonds: dict[str, str]) -> list[Position]:
    f = "depositaire_B_20260925.txt"
    lignes = (DATA / f).read_text(encoding="utf-8").splitlines()
    entete, details, fin = lignes[0], lignes[1:-1], lignes[-1]
    date = f"{entete[1:5]}-{entete[5:7]}-{entete[7:9]}"
    if date != DATE_ARRETE:
        tracer(f, 1, "date", entete[1:9], date, "BLOQUANT", "date ≠ arrêté")
    attendu = int(fin[1:])
    if attendu != len(details):
        tracer(f, len(lignes), "controle", fin, len(details), "BLOQUANT", "total de contrôle faux")
    out = []
    for n, ligne in enumerate(details, start=2):
        desk, isin, qte = ligne[1:4], ligne[4:16], ligne[16:31]
        q = parse_quantite(f, n, qte, "2 décimales implicites", lambda s: Decimal(int(s)) / 100)
        fonds = desk_vers_fonds.get(desk, "")
        if not fonds:
            tracer(f, n, "code_desk", desk, "", "BLOQUANT", "desk absent de la transco")
        out.append(Position("B", f, n, desk, fonds, isin, "", "B", q))
    return out


def lire_c(compte_vers_fonds: dict[str, str]) -> list[Position]:
    f = "depositaire_C_20260925.xml"
    ns = {"c": "urn:orsay:depositaire-c:positions:2026"}
    racine = ET.parse(DATA / f).getroot()
    d = racine.findtext("c:Entete/c:DateArrete", namespaces=ns)
    date = f"{d[4:]}-{d[2:4]}-{d[:2]}"
    if date != DATE_ARRETE:
        tracer(f, 0, "DateArrete", d, date, "BLOQUANT", "date ≠ arrêté")
    tracer(f, 0, "DateArrete", d, date, "NORMALISE", "JJMMAAAA → ISO")
    items = racine.findall("c:Positions/c:Item", ns)
    nb = int(racine.findtext("c:Entete/c:NbItems", namespaces=ns))
    if nb != len(items):
        tracer(f, 0, "NbItems", nb, len(items), "BLOQUANT", "total de contrôle faux")
    out = []
    for n, item in enumerate(items, start=1):
        p = {dim.get("prop"): (dim.text or "") for dim in item.findall("c:Dim", ns)}
        q = parse_quantite(
            f, n, p["QUANTITE"], "virgule décimale", lambda s: Decimal(s.replace(",", "."))
        )
        fonds = compte_vers_fonds.get(p["COMPTE"], "")
        if not fonds:
            tracer(f, n, "COMPTE", p["COMPTE"], "", "BLOQUANT", "compte absent de la transco")
        statut = "OK"
        if p["ACTION"] == "S":
            statut = "EXCLU"
            tracer(f, n, "ACTION", "S", "", "EXCLU", f"annulation (lot {p['LOT']})")
        isin = normaliser_isin(f, n, p["CODE_ISIN"])
        out.append(
            Position("C", f, n, p["COMPTE"], fonds, isin, p["LIBELLE_TITRE"], "C", q, statut)
        )
    return out


# --- contrôles transverses ---------------------------------------------------


def signaler(positions: list[Position], titres: dict[str, dict]):
    for p in positions:
        if p.statut == "EXCLU":
            continue
        if not isin_valide(p.isin):
            p.statut = "SIGNALE"
            tracer(p.fichier, p.ligne, "isin", p.isin, p.isin, "SIGNALE", "clé ISIN invalide")
        elif p.isin not in titres:
            p.statut = "SIGNALE"
            tracer(
                p.fichier, p.ligne, "isin", p.isin, p.isin, "SIGNALE", "ISIN absent du référentiel"
            )
        elif p.libelle and p.libelle.upper() != titres[p.isin]["libelle"].upper():
            tracer(
                p.fichier,
                p.ligne,
                "libelle",
                p.libelle,
                titres[p.isin]["libelle"],
                "SIGNALE",
                "libellé ≠ référentiel",
            )
        if not p.libelle and p.isin in titres:
            p.libelle = titres[p.isin]["libelle"]


def comparatif(positions: list[Position], titres: dict[str, dict]) -> list[dict]:
    """Agrège par (fonds, ISIN) : PMS par conservateur attendu, et quantité chez A/B/C."""
    agg: dict[tuple[str, str], dict] = {}
    for p in positions:
        if p.statut == "EXCLU":
            continue
        cle = (p.fonds, p.isin)
        r = agg.setdefault(
            cle,
            {
                "fonds": p.fonds,
                "isin": p.isin,
                "libelle": p.libelle,
                "type_titre": titres.get(p.isin, {}).get("type_titre", "?"),
                "PMS_A": Decimal(0),
                "PMS_B": Decimal(0),
                "PMS_C": Decimal(0),
                "A": Decimal(0),
                "B": Decimal(0),
                "C": Decimal(0),
            },
        )
        col = f"PMS_{p.conservateur}" if p.source == "PMS" else p.source
        r[col] += p.quantite
    for r in agg.values():
        r["ecart_A"] = r["A"] - r["PMS_A"]
        r["ecart_B"] = r["B"] - r["PMS_B"]
        r["ecart_C"] = r["C"] - r["PMS_C"]
        r["rapproche"] = all(r[k] == 0 for k in ("ecart_A", "ecart_B", "ecart_C"))
    return sorted(agg.values(), key=lambda r: (r["rapproche"], r["fonds"], r["isin"]))


def ecrire_csv(nom: str, lignes: list[dict]):
    if not lignes:
        return
    with open(OUT / nom, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(lignes[0]), delimiter=";")
        w.writeheader()
        w.writerows(lignes)


def main():
    OUT.mkdir(exist_ok=True)
    param = lire_parametrage()
    db = lire_db()
    nom_vers_code = {p["nom"]: p["code"] for p in db["portefeuilles"]}
    titres = {t["isin"]: t for t in db["titres"]}

    def transco(feuille, cle):
        return {str(r[cle]): nom_vers_code[r["entite"]] for r in param[feuille]}

    positions = (
        lire_pms(transco("TRANSCO_SOUS_COMPTES", "sous_compte"))
        + lire_a(transco("TRANSCO_DEPOSITAIRE_A", "compte"))
        + lire_b(transco("TRANSCO_DEPOSITAIRE_B", "code_desk"))
        + lire_c(transco("TRANSCO_DEPOSITAIRE_C", "compte"))
    )
    signaler(positions, titres)

    ecrire_csv("positions_unifiees.csv", [asdict(p) for p in positions])
    ecrire_csv("journal_nettoyage.csv", [asdict(t) for t in journal])
    comp = comparatif(positions, titres)
    ecrire_csv("comparatif.csv", comp)
    for nom, lignes in db.items():
        ecrire_csv(f"db_{nom}.csv", lignes)
    for nom, lignes in param.items():
        ecrire_csv(f"param_{nom}.csv", lignes)

    # résumé console
    from collections import Counter

    print(
        f"Positions lues : {len(positions)}  par source : {dict(Counter(p.source for p in positions))}"
    )
    print(f"Statuts : {dict(Counter(p.statut for p in positions))}")
    print(
        f"Journal : {len(journal)} entrées  par niveau : {dict(Counter(t.niveau for t in journal))}"
    )
    print("Règles :")
    for (niv, regle), n in sorted(Counter((t.niveau, t.regle) for t in journal).items()):
        print(f"  {niv:<10} {regle:<45} x{n}")
    ko = [r for r in comp if not r["rapproche"]]
    print(f"\nComparatif : {len(comp)} lignes (fonds, ISIN), {len(ko)} non rapprochées :")
    for r in ko:
        print(
            f"  {r['fonds']} {r['isin']} {r['libelle'][:28]:<28} "
            f"PMS A/B/C={r['PMS_A']}/{r['PMS_B']}/{r['PMS_C']}  dép A/B/C={r['A']}/{r['B']}/{r['C']}"
        )
    print(f"\nFichiers écrits dans {OUT.relative_to(ROOT)}/")


if __name__ == "__main__":
    main()
