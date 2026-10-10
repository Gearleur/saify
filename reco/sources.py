"""Lecture des sources (PMS, conservateurs A/B/C, paramétrage) en positions normalisées.

Chaque parseur rend des `Position` en `Decimal` et trace dans son journal toute transformation :
NORMALISE (sens inchangé), EXCLU (règle client), SIGNALE (douteux, conservé), BLOQUANT (fichier refusé).
"""

import csv
import io
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

SOUS_COMPTES_EXCLUS = {"P3G"}
NS_C = {"c": "urn:orsay:depositaire-c:positions:2026"}
PROPS_C = ("ACTION", "COMPTE", "CODE_ISIN", "LIBELLE_TITRE", "QUANTITE")


@dataclass(frozen=True)
class Position:
    source: str  # PMS, A, B, C
    entite: str  # nom d'entité du paramétrage, ex. ORSAY OBLIGATIONS
    isin: str
    conservateur: str  # A, B, C
    quantite: Decimal
    libelle: str
    fichier: str
    ligne: int  # numéro de ligne (n° d'Item pour le XML C)

    @property
    def ref(self) -> str:
        return f"{self.fichier}:{self.ligne}"


@dataclass(frozen=True)
class Anomalie:
    fichier: str
    ligne: int
    niveau: str  # NORMALISE, EXCLU, SIGNALE, BLOQUANT
    regle: str
    brut: str
    retenu: str
    motif: str


@dataclass
class Parametrage:
    sous_comptes: dict[str, str]
    comptes_a: dict[str, str]
    desks_b: dict[str, str]
    comptes_c: dict[str, str]
    types_ecarts: dict[str, tuple[str, str]]  # code -> (libellé, action attendue)


@dataclass
class Lecture:
    fichier: str
    positions: list[Position] = field(default_factory=list)
    journal: list[Anomalie] = field(default_factory=list)

    def noter(self, niveau: str, regle: str, ligne: int, brut, retenu, motif: str) -> None:
        self.journal.append(
            Anomalie(self.fichier, ligne, niveau, regle, str(brut), str(retenu), motif)
        )

    def ajouter(self, source: str, entite: str, isin: str, conservateur: str,
                quantite: Decimal, libelle: str, ligne: int) -> None:  # fmt: skip
        self.positions.append(
            Position(source, entite, isin, conservateur, quantite, libelle, self.fichier, ligne)
        )


def lire_parametrage(chemin: Path) -> Parametrage:
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)

    def table(nom: str) -> list[tuple]:
        lignes = list(wb[nom].iter_rows(values_only=True))
        return [ligne for ligne in lignes[1:] if ligne[0] is not None]

    def transco(nom: str) -> dict[str, str]:
        return {str(cle).strip(): str(entite).strip() for cle, entite, *_ in table(nom)}

    return Parametrage(
        sous_comptes=transco("TRANSCO_SOUS_COMPTES"),
        comptes_a=transco("TRANSCO_DEPOSITAIRE_A"),
        desks_b=transco("TRANSCO_DEPOSITAIRE_B"),
        comptes_c=transco("TRANSCO_DEPOSITAIRE_C"),
        types_ecarts={
            code: (libelle, action) for code, libelle, action, *_ in table("TYPES_ECARTS")
        },
    )


def code_conservateur(libelle: str) -> str | None:
    """CACEIS -> A, EUROCLEAR -> B, BNP/BP2S -> C (insensible à la casse) ; None si inconnu."""
    haut = libelle.upper()
    if "CACEIS" in haut:
        return "A"
    if "EUROCLEAR" in haut:
        return "B"
    if "BNP" in haut or "BP2S" in haut:
        return "C"
    return None


def isin_valide(isin: str) -> bool:
    """Format ISO 6166 et clé de Luhn (lettres converties en nombres, A=10 … Z=35)."""
    if not re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", isin):
        return False
    chiffres = "".join(str(int(c, 36)) for c in isin)
    total = 0
    for i, c in enumerate(reversed(chiffres)):
        n = int(c) * (2 if i % 2 else 1)
        total += n - 9 if n > 9 else n
    return total % 10 == 0


def _decimal(texte: str) -> Decimal | None:
    try:
        valeur = Decimal(texte)
    except InvalidOperation:
        return None
    return valeur if valeur.is_finite() else None


def _isin(lecture: Lecture, ligne: int, brut: str) -> str:
    isin = re.sub(r"\s", "", brut).upper()
    if isin != brut:
        lecture.noter("NORMALISE", "ISIN_FORMAT", ligne, brut, isin, "espaces et casse retirés")
    return isin


def _controler_date(lecture: Lecture, ligne: int, lue: date | None, attendue: date, brut) -> None:
    if lue != attendue:
        lecture.noter(
            "BLOQUANT", "DATE_ARRETE", ligne, brut, "",
            f"date d'arrêté {lue or 'introuvable'} ≠ {attendue:%d/%m/%Y}",
        )  # fmt: skip


def _texte(contenu: bytes) -> str:
    try:
        return contenu.decode("utf-8-sig")
    except UnicodeDecodeError:
        return contenu.decode("latin-1")


def lire_pms(nom: str, contenu: bytes, param: Parametrage, date_arrete: date) -> Lecture:
    lecture = Lecture(nom)
    lignes = csv.DictReader(io.StringIO(_texte(contenu), newline=""), delimiter=";")
    for num, ligne in enumerate(lignes, start=2):
        brut = ";".join(v or "" for v in ligne.values())
        try:
            lue = date.fromisoformat((ligne["date_arrete"] or "").strip())
        except ValueError:
            lue = None
        _controler_date(lecture, num, lue, date_arrete, brut)

        sous_compte = (ligne["sous_compte"] or "").strip()
        if not sous_compte:
            lecture.noter(
                "EXCLU",
                "SANS_SOUS_COMPTE",
                num,
                brut,
                "",
                "ligne sans sous-compte : anomalie de saisie",
            )
            continue
        if sous_compte in SOUS_COMPTES_EXCLUS:
            lecture.noter("EXCLU", "P3G", num, brut, "", "compte technique de passage exclu")
            continue
        if sous_compte not in param.sous_comptes:
            lecture.noter(
                "BLOQUANT",
                "TRANSCO",
                num,
                sous_compte,
                "",
                "sous-compte absent de TRANSCO_SOUS_COMPTES",
            )
            continue

        libelle_cons = (ligne["depositaire"] or "").strip()
        conservateur = code_conservateur(libelle_cons)
        if conservateur is None:
            lecture.noter(
                "BLOQUANT",
                "CONSERVATEUR_INCONNU",
                num,
                libelle_cons,
                "",
                "libellé de conservateur inconnu",
            )
            continue
        lecture.noter(
            "NORMALISE",
            "LIBELLE_CONSERVATEUR",
            num,
            libelle_cons,
            conservateur,
            "libellé → code conservateur",
        )

        quantite = _decimal((ligne["quantite"] or "").strip())
        if quantite is None:
            lecture.noter(
                "BLOQUANT",
                "QUANTITE_ILLISIBLE",
                num,
                ligne["quantite"],
                "",
                "quantité non numérique",
            )
            continue

        isin = _isin(lecture, num, ligne["isin"] or "")
        lecture.ajouter("PMS", param.sous_comptes[sous_compte], isin, conservateur, quantite,
                        (ligne["libelle"] or "").strip(), num)  # fmt: skip
    return lecture


def lire_depositaire_a(nom: str, contenu: bytes, param: Parametrage, date_arrete: date) -> Lecture:
    lecture = Lecture(nom)
    wb = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    lignes = list(wb.worksheets[0].iter_rows(values_only=True))

    titre = str(lignes[0][0] or "") if lignes else ""
    trouvee = re.search(r"(\d{2})/(\d{2})/(\d{4})", titre)
    lue = date(int(trouvee[3]), int(trouvee[2]), int(trouvee[1])) if trouvee else None
    _controler_date(lecture, 1, lue, date_arrete, titre)

    entete = next((i for i, ligne in enumerate(lignes) if ligne and ligne[0] == "Compte"), None)
    if entete is None:
        lecture.noter("BLOQUANT", "FORMAT", 0, "", "", "en-tête « Compte » introuvable")
        return lecture

    for num, (compte, isin, designation, quantite, *_) in enumerate(
        lignes[entete + 1 :], start=entete + 2
    ):
        if compte is None:
            continue
        compte = str(compte).strip()
        if compte not in param.comptes_a:
            lecture.noter(
                "BLOQUANT", "TRANSCO", num, compte, "", "compte absent de TRANSCO_DEPOSITAIRE_A"
            )
            continue
        brut = str(quantite)
        qte = _decimal(re.sub(r"\s", "", brut))
        if qte is None:
            lecture.noter("BLOQUANT", "QUANTITE_ILLISIBLE", num, brut, "", "quantité non numérique")
            continue
        if brut != str(qte):
            lecture.noter(
                "NORMALISE", "SEPARATEUR_MILLIERS", num, brut, qte, "quantité texte → nombre"
            )
        lecture.ajouter("A", param.comptes_a[compte], _isin(lecture, num, str(isin)), "A", qte,
                        str(designation or "").strip(), num)  # fmt: skip
    return lecture


def lire_depositaire_b(nom: str, contenu: bytes, param: Parametrage, date_arrete: date) -> Lecture:
    """Largeur fixe : H+date, D + desk(3) + ISIN(12) + quantité(15, 2 décimales implicites), T+nb."""
    lecture = Lecture(nom)
    lignes = _texte(contenu).splitlines()
    entete = next((ligne for ligne in lignes if ligne.startswith("H")), "")
    try:
        lue = date(int(entete[1:5]), int(entete[5:7]), int(entete[7:9]))
    except ValueError:
        lue = None
    _controler_date(lecture, 1, lue, date_arrete, entete)

    details = 0
    controle = None
    for num, ligne in enumerate(lignes, start=1):
        if ligne.startswith("T"):
            controle = ligne
        if not ligne.startswith("D"):
            continue
        details += 1
        desk, isin, brut = ligne[1:4], ligne[4:16], ligne[16:31]
        if len(ligne) != 31:
            lecture.noter("BLOQUANT", "FORMAT", num, ligne, "", "ligne de détail de longueur ≠ 31")
            continue
        if desk not in param.desks_b:
            lecture.noter(
                "BLOQUANT", "TRANSCO", num, desk, "", "desk absent de TRANSCO_DEPOSITAIRE_B"
            )
            continue
        if not brut.isdigit():
            lecture.noter("BLOQUANT", "QUANTITE_ILLISIBLE", num, brut, "", "quantité non numérique")
            continue
        qte = Decimal(brut).scaleb(-2)
        lecture.noter("NORMALISE", "DECIMALES_IMPLICITES", num, brut, qte, "2 décimales implicites")
        lecture.ajouter("B", param.desks_b[desk], _isin(lecture, num, isin), "B", qte, "", num)

    if controle is None or not controle[1:].isdigit() or int(controle[1:]) != details:
        lecture.noter("BLOQUANT", "TOTAL_CONTROLE", len(lignes), controle, details,
                      f"total de contrôle {controle!r} ≠ {details} lignes de détail")  # fmt: skip
    return lecture


def lire_depositaire_c(nom: str, contenu: bytes, param: Parametrage, date_arrete: date) -> Lecture:
    """XML à namespace ; les lots d'un même (compte, ISIN) sont additionnés au rapprochement."""
    lecture = Lecture(nom)
    racine = ET.fromstring(contenu)
    brut_date = racine.findtext(".//c:DateArrete", "", NS_C).strip()
    try:
        lue = date(int(brut_date[4:8]), int(brut_date[2:4]), int(brut_date[0:2]))
    except ValueError:
        lue = None
    _controler_date(lecture, 0, lue, date_arrete, brut_date)

    items = racine.findall(".//c:Item", NS_C)
    nb_items = racine.findtext(".//c:NbItems", "", NS_C).strip()
    if nb_items != str(len(items)):  # NbItems compte aussi les annulations
        lecture.noter("BLOQUANT", "TOTAL_CONTROLE", 0, nb_items, len(items),
                      f"NbItems {nb_items!r} ≠ {len(items)} Item lus")  # fmt: skip

    for num, item in enumerate(items, start=1):
        dims = {dim.get("prop"): (dim.text or "").strip() for dim in item.iterfind("c:Dim", NS_C)}
        manquants = [p for p in PROPS_C if p not in dims]
        if manquants:
            lecture.noter(
                "BLOQUANT",
                "CHAMP_MANQUANT",
                num,
                dims,
                "",
                f"Dim manquant : {', '.join(manquants)}",
            )
            continue
        if dims["ACTION"] == "S":
            lecture.noter(
                "EXCLU",
                "ANNULATION_S",
                num,
                dims,
                "",
                "item annulé (ACTION=S), ni ajouté ni soustrait",
            )
            continue
        if dims["COMPTE"] not in param.comptes_c:
            lecture.noter(
                "BLOQUANT",
                "TRANSCO",
                num,
                dims["COMPTE"],
                "",
                "compte absent de TRANSCO_DEPOSITAIRE_C",
            )
            continue
        brut = dims["QUANTITE"]
        qte = _decimal(brut.replace(",", "."))
        if qte is None:
            lecture.noter("BLOQUANT", "QUANTITE_ILLISIBLE", num, brut, "", "quantité non numérique")
            continue
        if "," in brut:
            lecture.noter("NORMALISE", "VIRGULE_DECIMALE", num, brut, qte, "virgule décimale")
        lecture.ajouter("C", param.comptes_c[dims["COMPTE"]], _isin(lecture, num, dims["CODE_ISIN"]), "C",
                        qte, dims["LIBELLE_TITRE"], num)  # fmt: skip
    return lecture
