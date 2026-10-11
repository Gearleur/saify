"""Paramétrage du client (parametrage.xlsx) : tables de transcodification et typologie des suspens."""

from contextlib import closing
from dataclasses import dataclass
from pathlib import Path

import openpyxl

REQUIRED_TYPES = ("QTE", "ABS_DEP", "ABS_INT", "DEP")


@dataclass
class ClientConfig:
    sub_accounts: dict[str, str]
    accounts_a: dict[str, str]
    desks_b: dict[str, str]
    accounts_c: dict[str, str]
    break_types: dict[str, tuple[str, str]]  # code -> (libellé, action attendue)


def read_client_config(path: Path) -> ClientConfig:
    """Lève ValueError, avec un message actionnable, si le paramétrage est incomplet ou contradictoire."""
    with closing(openpyxl.load_workbook(path, read_only=True, data_only=True)) as wb:
        return _client_config(wb)


def _client_config(wb) -> ClientConfig:

    def table(name: str) -> list[tuple]:
        if name not in wb.sheetnames:
            raise ValueError(f"feuille {name} absente")
        lines = list(wb[name].iter_rows(values_only=True))
        return [line for line in lines[1:] if line and line[0] is not None]

    def mapping(name: str) -> dict[str, str]:
        result: dict[str, str] = {}
        for key, entity, *_ in table(name):
            key, entity = str(key).strip(), str(entity).strip()
            if result.setdefault(key, entity) != entity:
                raise ValueError(
                    f"{name} : « {key} » associé à « {result[key]} » et à « {entity} »"
                )
        return result

    break_types = {
        str(code).strip(): (str(label or "").strip(), str(action or "").strip())
        for code, label, action, *_ in table("TYPES_ECARTS")
    }
    incomplete = [code for code, (label, action) in break_types.items() if not label or not action]
    if incomplete:
        raise ValueError(
            f"TYPES_ECARTS : libellé ou action attendue vide pour {', '.join(incomplete)}"
        )
    missing = [code for code in REQUIRED_TYPES if code not in break_types]
    if missing:
        raise ValueError(f"TYPES_ECARTS : code(s) {', '.join(missing)} absent(s)")
    return ClientConfig(
        sub_accounts=mapping("TRANSCO_SOUS_COMPTES"),
        accounts_a=mapping("TRANSCO_DEPOSITAIRE_A"),
        desks_b=mapping("TRANSCO_DEPOSITAIRE_B"),
        accounts_c=mapping("TRANSCO_DEPOSITAIRE_C"),
        break_types=break_types,
    )
