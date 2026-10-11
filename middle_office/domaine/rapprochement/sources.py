"""Lecture des sources (PMS, conservateurs A/B/C, paramétrage) en positions normalisées.

Chaque parseur rend des `Position` en `Decimal` et trace dans son journal toute transformation :
NORMALISE (sens inchangé), EXCLU (règle client), SIGNALE (douteux, conservé), BLOQUANT (fichier refusé).
Une donnée n'apparaît ni ne disparaît sans trace : une ligne vide est ignorée, une ligne partielle bloque.
"""

import csv
import io
import re
import xml.etree.ElementTree as ET
import zipfile
from contextlib import closing
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

SOUS_COMPTES_EXCLUS = {"P3G"}
NS_C = {"c": "urn:orsay:depositaire-c:positions:2026"}
PROPS_C = ("ACTION", "COMPTE", "CODE_ISIN", "LIBELLE_TITRE", "QUANTITE", "LOT")
ACTIONS_C = {"C", "S"}  # C : position, S : annulation
COLONNES_PMS = ("date_arrete", "sous_compte", "isin", "libelle", "quantite", "depositaire")
TYPES_REQUIS = ("QTE", "ABS_DEP", "ABS_INT", "DEP")
# Ce qu'un fichier mal formé peut lever à la lecture (vérifié sur des fichiers vides, binaires,
# tronqués, en UTF-16). Le reste est un bug de l'outil : il doit remonter, pas devenir « fichier illisible ».
ERREURS_DE_FORMAT = (
    ValueError,
    KeyError,
    IndexError,
    EOFError,
    csv.Error,
    ET.ParseError,
    zipfile.BadZipFile,
    InvalidFileException,
)

# Ce que l'opérateur doit faire quand un fichier est refusé : un blocage doit être actionnable.
REMEDES = {
    "DATE_ARRETE": "Redemander le fichier arrêté à la bonne date à son émetteur, puis relancer",
    "TOTAL_CONTROLE": "Fichier tronqué ou altéré : le redemander au conservateur, puis relancer",
    "QUANTITE_ILLISIBLE": "Corriger la quantité de la ligne citée à la source, puis relancer",
    "CONSERVATEUR_INCONNU": "Corriger le libellé de dépositaire « {brut} » dans le PMS, puis relancer",
    "LIGNE_INCOMPLETE": "Compléter ou supprimer la ligne citée à la source, puis relancer",
    "TYPE_LIGNE_INCONNU": "Type d'enregistrement inconnu : vérifier le format avec le conservateur",
    "CHAMP_MANQUANT": "Item incomplet : redemander le fichier au conservateur, puis relancer",
    "DOUBLON": "Ligne en double : faire confirmer la position par son émetteur, supprimer le doublon, puis relancer",
    "ACTION_INCONNUE": "Code ACTION inconnu : vérifier sa signification avec le conservateur avant de relancer",
    "FORMAT": "Le format du fichier a changé : le vérifier avec son émetteur avant de relancer",
    "FORMAT_ILLISIBLE": "Redéposer le fichier d'origine, sans l'ouvrir ni le réenregistrer",
    "AUCUNE_LIGNE": "Aucune position lue : vérifier qu'il s'agit du bon fichier et du bon format",
    "FICHIER_MANQUANT": "Déposer le fichier manquant, puis relancer",
    "FICHIERS_MULTIPLES": "Ne déposer qu'un fichier par source, puis relancer",
    "PARAMETRAGE": "Corriger parametrage.xlsx comme indiqué, puis relancer",
}


@dataclass(frozen=True)
class Position:
    source: str  # PMS, A, B, C
    entite: str  # nom d'entité du paramétrage, ex. ORSAY OBLIGATIONS
    isin: str
    conservateur: str  # A, B, C
    quantite: Decimal
    libelle: str
    fichier: str
    ligne: int  # numéro de ligne (n° d'Item pour le XML C)

    @property
    def ref(self) -> str:
        return f"{self.fichier}:{self.ligne}"


@dataclass(frozen=True)
class Anomalie:
    fichier: str
    ligne: int
    niveau: str  # NORMALISE, EXCLU, SIGNALE, BLOQUANT
    regle: str
    brut: str
    retenu: str
    motif: str
    remede: str = ""  # renseigné pour tout BLOQUANT : quoi faire pour débloquer


def anomalie(
    fichier: str,
    ligne: int,
    niveau: str,
    regle: str,
    brut,
    retenu,
    motif: str,
    remede: str | None = None,
) -> Anomalie:
    if remede is None and niveau == "BLOQUANT":
        remede = REMEDES[regle].format(brut=brut)
    return Anomalie(fichier, ligne, niveau, regle, str(brut), str(retenu), motif, remede or "")


@dataclass
class Parametrage:
    sous_comptes: dict[str, str]
    comptes_a: dict[str, str]
    desks_b: dict[str, str]
    comptes_c: dict[str, str]
    types_ecarts: dict[str, tuple[str, str]]  # code -> (libellé, action attendue)


@dataclass
class Lecture:
    fichier: str
    positions: list[Position] = field(default_factory=list)
    journal: list[Anomalie] = field(default_factory=list)
    _vues: set[tuple] = field(default_factory=set)

    def noter(
        self,
        niveau: str,
        regle: str,
        ligne: int,
        brut,
        retenu,
        motif: str,
        remede: str | None = None,
    ) -> None:
        self.journal.append(
            anomalie(self.fichier, ligne, niveau, regle, brut, retenu, motif, remede)
        )

    def transco_absente(self, ligne: int, valeur: str, table: str) -> None:
        self.noter(
            "BLOQUANT",
            "TRANSCO",
            ligne,
            valeur,
            "",
            f"« {valeur} » absent de {table}",
            f"Ajouter « {valeur} » à la table {table} de parametrage.xlsx, puis relancer",
        )

    def ajouter(
        self,
        source: str,
        entite: str,
        isin: str,
        conservateur: str,
        quantite: Decimal,
        libelle: str,
        ligne: int,
        compte: str,
    ) -> None:
        """`compte` : identifiant de ligne côté source (pour C : compte et numéro de lot). Deux lignes
        au même (compte, ISIN) bloquent : additionnées, elles pourraient égaler le PMS et masquer
        un écart ; en garder une, ce serait deviner laquelle est la bonne."""
        cle = (compte, isin, conservateur)
        if cle in self._vues:
            self.noter(
                "BLOQUANT",
                "DOUBLON",
                ligne,
                f"{compte} {isin}",
                "",
                "ligne en double (même compte et ISIN) : laquelle est la bonne ?",
            )
            return
        self._vues.add(cle)
        if quantite < 0 or quantite != quantite.to_integral_value():
            self.noter(
                "SIGNALE",
                "QUANTITE_SUSPECTE",
                ligne,
                quantite,
                quantite,
                "quantité négative ou non entière : à confirmer",
            )
        self.positions.append(
            Position(source, entite, isin, conservateur, quantite, libelle, self.fichier, ligne)
        )


def lire_parametrage(chemin: Path) -> Parametrage:
    """Lève ValueError, avec un message actionnable, si le paramétrage est incomplet ou contradictoire."""
    with closing(openpyxl.load_workbook(chemin, read_only=True, data_only=True)) as wb:
        return _parametrage(wb)


def _parametrage(wb) -> Parametrage:

    def table(nom: str) -> list[tuple]:
        if nom not in wb.sheetnames:
            raise ValueError(f"feuille {nom} absente")
        lignes = list(wb[nom].iter_rows(values_only=True))
        return [ligne for ligne in lignes[1:] if ligne and ligne[0] is not None]

    def transco(nom: str) -> dict[str, str]:
        resultat: dict[str, str] = {}
        for cle, entite, *_ in table(nom):
            cle, entite = str(cle).strip(), str(entite).strip()
            if resultat.setdefault(cle, entite) != entite:
                raise ValueError(
                    f"{nom} : « {cle} » associé à « {resultat[cle]} » et à « {entite} »"
                )
        return resultat

    types_ecarts = {
        str(code).strip(): (str(libelle or "").strip(), str(action or "").strip())
        for code, libelle, action, *_ in table("TYPES_ECARTS")
    }
    incomplets = [
        code for code, (libelle, action) in types_ecarts.items() if not libelle or not action
    ]
    if incomplets:
        raise ValueError(
            f"TYPES_ECARTS : libellé ou action attendue vide pour {', '.join(incomplets)}"
        )
    manquants = [code for code in TYPES_REQUIS if code not in types_ecarts]
    if manquants:
        raise ValueError(f"TYPES_ECARTS : code(s) {', '.join(manquants)} absent(s)")
    return Parametrage(
        sous_comptes=transco("TRANSCO_SOUS_COMPTES"),
        comptes_a=transco("TRANSCO_DEPOSITAIRE_A"),
        desks_b=transco("TRANSCO_DEPOSITAIRE_B"),
        comptes_c=transco("TRANSCO_DEPOSITAIRE_C"),
        types_ecarts=types_ecarts,
    )


def code_conservateur(libelle: str) -> str | None:
    """CACEIS -> A, EUROCLEAR -> B, BNP/BP2S -> C (insensible à la casse) ; None si inconnu."""
    haut = libelle.upper()
    if "CACEIS" in haut:
        return "A"
    if "EUROCLEAR" in haut:
        return "B"
    if "BNP" in haut or "BP2S" in haut:
        return "C"
    return None


def isin_valide(isin: str) -> bool:
    """Format ISO 6166 et clé de Luhn (lettres converties en nombres, A=10 … Z=35)."""
    if not re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", isin):
        return False
    chiffres = "".join(str(int(c, 36)) for c in isin)
    total = 0
    for i, c in enumerate(reversed(chiffres)):
        n = int(c) * (2 if i % 2 else 1)
        total += n - 9 if n > 9 else n
    return total % 10 == 0


def _decimal(texte: str) -> Decimal | None:
    try:
        valeur = Decimal(texte)
    except InvalidOperation:
        return None
    return valeur if valeur.is_finite() else None


def _separateur_ambigu(lecture: Lecture, ligne: int, brut: str) -> None:
    """« 10.000 » se lit 10 mais veut sans doute dire 10 000 : la valeur est gardée, un humain confirme."""
    if re.search(r"\d[.,]\d{3}$", brut.strip()):
        lecture.noter(
            "SIGNALE",
            "QUANTITE_SUSPECTE",
            ligne,
            brut,
            brut,
            "séparateur suivi de 3 chiffres : séparateur de milliers probable",
        )


def _vide(valeur) -> bool:
    return valeur is None or str(valeur).strip() == ""


def _isin(lecture: Lecture, ligne: int, brut: str) -> str:
    isin = re.sub(r"\s", "", brut).upper()
    if isin != brut:
        lecture.noter("NORMALISE", "ISIN_FORMAT", ligne, brut, isin, "espaces et casse retirés")
    return isin


def _controler_date(
    lecture: Lecture, ligne: int, lue: date | None, attendue: date | None, brut
) -> None:
    """`attendue` None : la date de référence n'a pas pu être établie, ce blocage est déjà tracé."""
    if attendue is not None and lue != attendue:
        lecture.noter(
            "BLOQUANT",
            "DATE_ARRETE",
            ligne,
            brut,
            "",
            f"date d'arrêté {lue or 'introuvable'} ≠ {attendue:%d/%m/%Y}",
        )


def dates_pms(contenu: bytes) -> set[str]:
    """Dates d'arrêté distinctes du PMS : la source interne fixe la date du rapprochement."""
    texte = _texte(Lecture(""), contenu)
    lignes = csv.DictReader(io.StringIO(texte, newline=""), delimiter=";")
    return {(ligne.get("date_arrete") or "").strip() for ligne in lignes}


def _texte(lecture: Lecture, contenu: bytes) -> str:
    try:
        return contenu.decode("utf-8-sig")
    except UnicodeDecodeError:
        lecture.noter(
            "NORMALISE", "ENCODAGE", 0, "non UTF-8", "latin-1", "fichier décodé en latin-1"
        )
        return contenu.decode("latin-1")


# `date_arrete` vaut None quand le PMS n'a pas permis de l'établir : ce blocage est déjà tracé, les
# lecteurs lisent quand même pour signaler tout le reste en une fois.
def lire_pms(nom: str, contenu: bytes, param: Parametrage, date_arrete: date | None) -> Lecture:
    lecture = Lecture(nom)
    lignes = csv.DictReader(io.StringIO(_texte(lecture, contenu), newline=""), delimiter=";")
    absentes = [c for c in COLONNES_PMS if c not in (lignes.fieldnames or [])]
    if absentes:
        lecture.noter(
            "BLOQUANT",
            "FORMAT",
            1,
            ";".join(lignes.fieldnames or []),
            "",
            f"colonne(s) absente(s) : {', '.join(absentes)}",
        )
        return lecture

    for num, ligne in enumerate(lignes, start=2):
        brut = ";".join(str(v or "") for v in ligne.values())
        try:
            lue = date.fromisoformat((ligne["date_arrete"] or "").strip())
        except ValueError:
            lue = None
        _controler_date(lecture, num, lue, date_arrete, brut)

        sous_compte = (ligne["sous_compte"] or "").strip()
        if not sous_compte:
            lecture.noter(
                "EXCLU",
                "SANS_SOUS_COMPTE",
                num,
                brut,
                "",
                "ligne sans sous-compte : anomalie de saisie",
            )
            continue
        if sous_compte in SOUS_COMPTES_EXCLUS:
            lecture.noter("EXCLU", "P3G", num, brut, "", "compte technique de passage exclu")
            continue
        if sous_compte not in param.sous_comptes:
            lecture.transco_absente(num, sous_compte, "TRANSCO_SOUS_COMPTES")
            continue
        if _vide(ligne["isin"]):
            lecture.noter("BLOQUANT", "LIGNE_INCOMPLETE", num, brut, "", "ISIN vide")
            continue

        libelle_cons = (ligne["depositaire"] or "").strip()
        conservateur = code_conservateur(libelle_cons)
        if conservateur is None:
            lecture.noter(
                "BLOQUANT",
                "CONSERVATEUR_INCONNU",
                num,
                libelle_cons,
                "",
                "libellé de conservateur inconnu",
            )
            continue
        lecture.noter(
            "NORMALISE",
            "LIBELLE_CONSERVATEUR",
            num,
            libelle_cons,
            conservateur,
            "libellé → code conservateur",
        )

        quantite = _decimal((ligne["quantite"] or "").strip())
        _separateur_ambigu(lecture, num, ligne["quantite"] or "")
        if quantite is None:
            lecture.noter(
                "BLOQUANT",
                "QUANTITE_ILLISIBLE",
                num,
                ligne["quantite"],
                "",
                "quantité non numérique",
            )
            continue

        isin = _isin(lecture, num, ligne["isin"])
        lecture.ajouter(
            "PMS",
            param.sous_comptes[sous_compte],
            isin,
            conservateur,
            quantite,
            (ligne["libelle"] or "").strip(),
            num,
            compte=sous_compte,
        )
    return lecture


def lire_depositaire_a(
    nom: str, contenu: bytes, param: Parametrage, date_arrete: date | None
) -> Lecture:
    lecture = Lecture(nom)
    with closing(openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)) as wb:
        lignes = list(wb.worksheets[0].iter_rows(values_only=True))

    titre = str(lignes[0][0] or "") if lignes else ""
    trouvee = re.search(r"(\d{2})/(\d{2})/(\d{4})", titre)
    lue = date(int(trouvee[3]), int(trouvee[2]), int(trouvee[1])) if trouvee else None
    _controler_date(lecture, 1, lue, date_arrete, titre)

    entete = next((i for i, ligne in enumerate(lignes) if ligne and ligne[0] == "Compte"), None)
    if entete is None:
        lecture.noter("BLOQUANT", "FORMAT", 0, "", "", "en-tête « Compte » introuvable")
        return lecture

    for num, cellules in enumerate(lignes[entete + 1 :], start=entete + 2):
        compte, isin, designation, quantite = (list(cellules) + [None] * 4)[:4]
        if all(_vide(v) for v in cellules):
            continue
        if _vide(compte) or _vide(isin) or _vide(quantite):
            lecture.noter(
                "BLOQUANT", "LIGNE_INCOMPLETE", num, cellules, "", "compte, ISIN ou quantité vide"
            )
            continue
        compte = str(compte).strip()
        if compte not in param.comptes_a:
            lecture.transco_absente(num, compte, "TRANSCO_DEPOSITAIRE_A")
            continue
        brut = str(quantite)
        qte = _decimal(re.sub(r"\s", "", brut))
        _separateur_ambigu(lecture, num, brut)
        if qte is None:
            lecture.noter("BLOQUANT", "QUANTITE_ILLISIBLE", num, brut, "", "quantité non numérique")
            continue
        if brut != str(qte):
            lecture.noter(
                "NORMALISE", "SEPARATEUR_MILLIERS", num, brut, qte, "quantité texte → nombre"
            )
        lecture.ajouter(
            "A",
            param.comptes_a[compte],
            _isin(lecture, num, str(isin)),
            "A",
            qte,
            str(designation or "").strip(),
            num,
            compte=compte,
        )
    return lecture


def lire_depositaire_b(
    nom: str, contenu: bytes, param: Parametrage, date_arrete: date | None
) -> Lecture:
    """Largeur fixe : H+date, D + desk(3) + ISIN(12) + quantité(15, 2 décimales implicites), T+nb."""
    lecture = Lecture(nom)
    lignes = _texte(lecture, contenu).splitlines()
    entete = next((ligne for ligne in lignes if ligne.startswith("H")), "")
    try:
        lue = date(int(entete[1:5]), int(entete[5:7]), int(entete[7:9]))
    except ValueError:
        lue = None
    _controler_date(lecture, 1, lue, date_arrete, entete)

    details = 0
    controle = None
    for num, ligne in enumerate(lignes, start=1):
        if not ligne.strip():
            continue
        if ligne[0] not in "HDT":
            lecture.noter(
                "BLOQUANT",
                "TYPE_LIGNE_INCONNU",
                num,
                ligne,
                "",
                f"enregistrement de type « {ligne[0]} »",
            )
            continue
        if ligne.startswith("T"):
            controle = ligne
        if not ligne.startswith("D"):
            continue
        details += 1
        desk, isin, brut = ligne[1:4], ligne[4:16], ligne[16:31]
        if len(ligne) != 31:
            lecture.noter("BLOQUANT", "FORMAT", num, ligne, "", "ligne de détail de longueur ≠ 31")
            continue
        if desk not in param.desks_b:
            lecture.transco_absente(num, desk, "TRANSCO_DEPOSITAIRE_B")
            continue
        if not brut.isdigit():
            lecture.noter("BLOQUANT", "QUANTITE_ILLISIBLE", num, brut, "", "quantité non numérique")
            continue
        qte = Decimal(brut).scaleb(-2)
        lecture.noter("NORMALISE", "DECIMALES_IMPLICITES", num, brut, qte, "2 décimales implicites")
        lecture.ajouter(
            "B", param.desks_b[desk], _isin(lecture, num, isin), "B", qte, "", num, compte=desk
        )

    if controle is None or not controle[1:].isdigit() or int(controle[1:]) != details:
        lecture.noter(
            "BLOQUANT",
            "TOTAL_CONTROLE",
            len(lignes),
            controle,
            details,
            f"total de contrôle {controle!r} ≠ {details} lignes de détail",
        )
    return lecture


def _dtd_present(contenu: bytes) -> bool:
    """Cherche aussi dans la version UTF-16, qu'expat accepterait avec un BOM."""
    textes = [contenu.upper()]
    if contenu[:2] in (b"\xff\xfe", b"\xfe\xff"):
        textes.append(contenu.decode("utf-16", errors="replace").upper().encode())
    return any(b"<!DOCTYPE" in t or b"<!ENTITY" in t for t in textes)


def lire_depositaire_c(
    nom: str, contenu: bytes, param: Parametrage, date_arrete: date | None
) -> Lecture:
    """XML à namespace ; les lots distincts d'un même (compte, ISIN) sont additionnés, un lot répété bloque."""
    lecture = Lecture(nom)
    if _dtd_present(contenu):  # avant tout parsing : entités externes, « billion laughs »
        lecture.noter(
            "BLOQUANT",
            "FORMAT",
            0,
            "",
            "",
            "DOCTYPE ou ENTITY interdit dans le XML conservateur : fichier refusé par sécurité",
        )
        return lecture
    racine = ET.fromstring(contenu)
    brut_date = racine.findtext(".//c:DateArrete", "", NS_C).strip()
    try:
        lue = date(int(brut_date[4:8]), int(brut_date[2:4]), int(brut_date[0:2]))
    except ValueError:
        lue = None
    _controler_date(lecture, 0, lue, date_arrete, brut_date)

    items = racine.findall(".//c:Item", NS_C)
    nb_items = racine.findtext(".//c:NbItems", "", NS_C).strip()
    if nb_items != str(len(items)):  # NbItems compte aussi les annulations
        lecture.noter(
            "BLOQUANT",
            "TOTAL_CONTROLE",
            0,
            nb_items,
            len(items),
            f"NbItems {nb_items!r} ≠ {len(items)} Item lus",
        )

    for num, item in enumerate(items, start=1):
        dims = {dim.get("prop"): (dim.text or "").strip() for dim in item.iterfind("c:Dim", NS_C)}
        manquants = [p for p in PROPS_C if p not in dims or (p != "LIBELLE_TITRE" and not dims[p])]
        if manquants:
            lecture.noter(
                "BLOQUANT",
                "CHAMP_MANQUANT",
                num,
                dims,
                "",
                f"Dim manquant ou vide : {', '.join(manquants)}",
            )
            continue
        if dims["ACTION"] == "S":
            lecture.noter(
                "EXCLU",
                "ANNULATION_S",
                num,
                dims,
                "",
                "item annulé (ACTION=S), ni ajouté ni soustrait",
            )
            continue
        if dims["ACTION"] not in ACTIONS_C:
            # Ni additionné (il pourrait combler un écart) ni écarté (il pourrait en créer un) : refusé.
            lecture.noter(
                "BLOQUANT",
                "ACTION_INCONNUE",
                num,
                dims["ACTION"],
                "",
                f"ACTION « {dims['ACTION']} » inconnue (attendu : C ou S)",
            )
            continue
        if dims["COMPTE"] not in param.comptes_c:
            lecture.transco_absente(num, dims["COMPTE"], "TRANSCO_DEPOSITAIRE_C")
            continue
        brut = dims["QUANTITE"]
        qte = _decimal(brut.replace(",", "."))
        _separateur_ambigu(lecture, num, brut)
        if qte is None:
            lecture.noter("BLOQUANT", "QUANTITE_ILLISIBLE", num, brut, "", "quantité non numérique")
            continue
        if "," in brut:
            lecture.noter("NORMALISE", "VIRGULE_DECIMALE", num, brut, qte, "virgule décimale")
        lecture.ajouter(
            "C",
            param.comptes_c[dims["COMPTE"]],
            _isin(lecture, num, dims["CODE_ISIN"]),
            "C",
            qte,
            dims["LIBELLE_TITRE"],
            num,
            compte=f"{dims['COMPTE']} lot {dims['LOT']}",
        )
    return lecture
