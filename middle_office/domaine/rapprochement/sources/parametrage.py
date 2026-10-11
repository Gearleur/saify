"""Paramétrage du client (parametrage.xlsx) : tables de transcodification et typologie des suspens."""

from contextlib import closing
from dataclasses import dataclass
from pathlib import Path

import openpyxl

TYPES_REQUIS = ("QTE", "ABS_DEP", "ABS_INT", "DEP")


@dataclass
class Parametrage:
    sous_comptes: dict[str, str]
    comptes_a: dict[str, str]
    desks_b: dict[str, str]
    comptes_c: dict[str, str]
    types_ecarts: dict[str, tuple[str, str]]  # code -> (libellé, action attendue)


def lire_parametrage(chemin: Path) -> Parametrage:
    """Lève ValueError, avec un message actionnable, si le paramétrage est incomplet ou contradictoire."""
    with closing(openpyxl.load_workbook(chemin, read_only=True, data_only=True)) as wb:
        return _parametrage(wb)


def _parametrage(wb) -> Parametrage:

    def table(nom: str) -> list[tuple]:
        if nom not in wb.sheetnames:
            raise ValueError(f"feuille {nom} absente")
        lignes = list(wb[nom].iter_rows(values_only=True))
        return [ligne for ligne in lignes[1:] if ligne and ligne[0] is not None]

    def transco(nom: str) -> dict[str, str]:
        resultat: dict[str, str] = {}
        for cle, entite, *_ in table(nom):
            cle, entite = str(cle).strip(), str(entite).strip()
            if resultat.setdefault(cle, entite) != entite:
                raise ValueError(
                    f"{nom} : « {cle} » associé à « {resultat[cle]} » et à « {entite} »"
                )
        return resultat

    types_ecarts = {
        str(code).strip(): (str(libelle or "").strip(), str(action or "").strip())
        for code, libelle, action, *_ in table("TYPES_ECARTS")
    }
    incomplets = [
        code for code, (libelle, action) in types_ecarts.items() if not libelle or not action
    ]
    if incomplets:
        raise ValueError(
            f"TYPES_ECARTS : libellé ou action attendue vide pour {', '.join(incomplets)}"
        )
    manquants = [code for code in TYPES_REQUIS if code not in types_ecarts]
    if manquants:
        raise ValueError(f"TYPES_ECARTS : code(s) {', '.join(manquants)} absent(s)")
    return Parametrage(
        sous_comptes=transco("TRANSCO_SOUS_COMPTES"),
        comptes_a=transco("TRANSCO_DEPOSITAIRE_A"),
        desks_b=transco("TRANSCO_DEPOSITAIRE_B"),
        comptes_c=transco("TRANSCO_DEPOSITAIRE_C"),
        types_ecarts=types_ecarts,
    )
